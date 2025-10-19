# scripts/generate_phase_summary.py
# Builds results/Phase0_to_Phase6_Summary_with_Lessons.docx from your Phase 0–6 artifacts

from pathlib import Path
from datetime import datetime
import json, math
import pandas as pd

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

RESULTS = Path("results")
RESULTS.mkdir(parents=True, exist_ok=True)

# ---------- helpers ----------
def load_json(p: Path):
    try:
        with open(p, "r") as f:
            return json.load(f)
    except Exception:
        return {}

def pct(x):
    if x is None: return ""
    try:
        x = float(x)
        return f"{x*100:.2f}%" if x <= 1 else f"{x:.2f}%"
    except Exception:
        return str(x)

def row_from(d: dict):
    return [
        d.get("total_trades", d.get("Total trades","")),
        pct(d.get("win_rate", d.get("Win rate",""))),
        d.get("profit_factor", d.get("PF","")),
        d.get("avg_r", d.get("Average R","")),
        d.get("expected_value_r", d.get("Expected Val. R", d.get("Expected Value",""))),
        d.get("max_drawdown_r", d.get("Max drawdown","")),
    ]

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i,h in enumerate(headers): t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text = "" if val is None else str(val)
    return t

# ---------- locate inputs ----------
rules_summary = RESULTS / "backtest_summary.json"                # rules only
ml_summary    = RESULTS / "phase6_ml_summary.json"               # ml only
hyb_summary   = RESULTS / "phase6_hybrid_summary.json"           # hybrid
hyb_trades    = RESULTS / "phase6_hybrid_trades.csv"             # for risk metrics
hyb_sessions  = RESULTS / "phase6_hybrid_session_summary.csv"    # optional appendix

rules = load_json(rules_summary) if rules_summary.exists() else {}
ml    = load_json(ml_summary)    if ml_summary.exists()    else {}
hyb   = load_json(hyb_summary)   if hyb_summary.exists()   else {}

# ---------- build doc ----------
doc = Document()

# title
h = doc.add_heading("TCN_1.0, Phase 0 to Phase 6, Summary and Lessons", level=0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph()
p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").italic = True

# overview
doc.add_heading("1. Overview", level=1)
doc.add_paragraph(
    "Phases 0 to 6 completed, data prepared, rules backtests built with EV, "
    "TCN predictions integrated, and a hybrid backtest produced stable results under a common cost model."
)

# phases list
doc.add_heading("2. Phases 0 to 6, highlights", level=1)
phases = [
    ("Phase 0", "Environment and repo, virtual env, dependencies, results layout"),
    ("Phase 1", "Data cleaning, timestamp alignment, synthetic timestamps when missing"),
    ("Phase 2", "Feature engineering, ATR, EMA, ADX, OBV, sessions, regime flags"),
    ("Phase 3", "Dataset assembly, model ready frame, metadata saved"),
    ("Phase 4", "Prototype TCN, scaler persisted to models"),
    ("Phase 5", "Rules backtests with EV, grid and filter sweeps, diagnostics"),
    ("Phase 6", "ML predictions and Hybrid backtest, probability gating on rules"),
]
t = doc.add_table(rows=1, cols=3); t.style="Table Grid"
t.rows[0].cells[0].text="Phase"; t.rows[0].cells[1].text="Focus"; t.rows[0].cells[2].text="Notes"
for a,b in phases:
    r = t.add_row().cells
    r[0].text=a; r[1].text=b; r[2].text=""

# results table
doc.add_heading("3. Results summary", level=1)
hdr = ["System","Trades","Win Rate","Profit Factor","Avg R","EV (R)","Max DD (R)"]
tbl = doc.add_table(rows=1, cols=len(hdr)); tbl.style="Table Grid"
for i,hc in enumerate(hdr): tbl.rows[0].cells[i].text = hc
for sys, data in [
    ("Rules only", row_from(rules)),
    ("ML only",    row_from(ml)),
    ("Hybrid",     row_from(hyb)),
]:
    cells = tbl.add_row().cells
    cells[0].text = sys
    for i,v in enumerate(data, start=1):
        cells[i].text = "" if v is None else str(v)

# risk metrics from hybrid trades
doc.add_heading("3.1 Risk metrics, Hybrid", level=2)
if hyb_trades.exists():
    d = pd.read_csv(hyb_trades)
    r = d["R"] if "R" in d.columns else d.get("R_net", pd.Series([0.0]*len(d)))
    cum = r.cumsum()
    dd = (cum.cummax() - cum).max()
    final_cum_r = float(cum.iloc[-1]) if len(cum) else 0.0
    sharpe = float((r.mean() / r.std()) * 252**0.5) if r.std() else float("nan")
    neg_std = r[r < 0].std()
    sortino = float((r.mean() / neg_std) * 252**0.5) if neg_std else float("nan")
    mar = float(final_cum_r / dd) if dd and not math.isclose(dd, 0.0) else float("nan")
    add_table(
        doc,
        ["Metric","Value"],
        [
            ["Final Cumulative R", f"{final_cum_r:.2f}"],
            ["Sharpe Ratio", f"{sharpe:.2f}"],
            ["Sortino Ratio", f"{sortino:.2f}"],
            ["Max Drawdown (R)", f"{dd:.2f}"],
            ["MAR Ratio", f"{mar:.2f}"],
        ],
    )
else:
    doc.add_paragraph("Hybrid trades file not found, risk metrics skipped.")

# best config
doc.add_heading("4. Best configuration", level=1)
best = {
    "Combine mode": "ml_gate_rules",
    "ML threshold": "0.55",
    "Horizon bars": hyb.get("horizon_bars", hyb.get("Horizon Bars","20")),
    "Spread pips":  hyb.get("spread_pips", hyb.get("Spread (pips)","")),
    "Commission USD": hyb.get("commission_usd", hyb.get("Commission (USD)","")),
}
for k,v in best.items():
    p = doc.add_paragraph(); r1 = p.add_run(f"{k}: "); r1.bold=True; p.add_run(str(v))

# lessons
doc.add_heading("5. Lessons learned", level=1)
for s in [
    "Align features to the scaler list and keep preprocessing identical for training and inference",
    "Normalize timestamps to UTC and round to bar size before merges",
    "Use ML to gate rules for precision, avoid permissive either mode",
    "Compare EV and drawdown together, not Profit Factor alone",
    "Keep cost model and horizon constant across runs for fair comparisons",
]:
    doc.add_paragraph(f"• {s}")

# reproducibility and Git notes
doc.add_heading("6. Repro and Git notes", level=1)
doc.add_paragraph("Winning backtest command, PowerShell:")
doc.add_paragraph(
    "python scripts\\phase6_step4_hybrid_backtest.py "
    "--features data\\m15_features.parquet "
    "--predictions results\\phase6_predictions.csv "
    "--rules_trades results\\phase5_rules_trades.csv "
    "--combine_mode ml_gate_rules "
    "--ml_threshold 0.55 "
    "--out_dir results\\final_ml_gate_rules_th_055"
)
doc.add_paragraph("Push commits and tags, PowerShell:")
doc.add_paragraph("git push ; git push --tags")
doc.add_paragraph("Ignore generated artifacts in Git:")
doc.add_paragraph("add to .gitignore -> results/*.csv, results/**/*.csv, results/*.json, results/**/*.json, models/*.pt, models/*.pkl")

# sessions appendix
if hyb_sessions.exists():
    doc.add_heading("Appendix, session summary, hybrid", level=2)
    sess = pd.read_csv(hyb_sessions)
    cols = [c for c in sess.columns if c.lower() in {"session","trades","win_rate","profit_factor","avg_r","ev_r","sum_r"}]
    if cols:
        tt = doc.add_table(rows=1, cols=len(cols)); tt.style="Table Grid"
        for i,hc in enumerate(cols): tt.rows[0].cells[i].text = hc
        for _,row in sess[cols].iterrows():
            cells = tt.add_row().cells
            for i,hc in enumerate(cols):
                val = row[hc]
                cells[i].text = "" if pd.isna(val) else str(val)

out = RESULTS / "Phase0_to_Phase6_Summary_with_Lessons.docx"
doc.save(out)
print(f"Saved: {out}")
